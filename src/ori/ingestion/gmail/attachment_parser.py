from __future__ import annotations

import io
import logging

import pytesseract
from docx import Document
from PIL import Image
from pypdf import PdfReader

# a bigger attachment isn't worth the extraction time/memory for a
# personal-assistant use case (a scanned 200-page PDF, a multi-GB video
# misnamed with a doc extension) - skip rather than let one huge file
# stall a sync. Chosen well above any real personal document (a few MB)
# but well below anything pathological.
_MAX_ATTACHMENT_BYTES = 15_000_000

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_PDF_MIME = "application/pdf"
_IMAGE_MIMES = {"image/png", "image/jpeg", "image/jpg", "image/tiff", "image/bmp", "image/webp"}
_IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp")


def extract_attachment_text(
    filename: str, mime_type: str, data: bytes, *, logger: logging.Logger | None = None,
) -> str:
    """best-effort text extraction from one email attachment (.docx, .pdf,
    or an image via OCR). Never raises to the caller - matches this
    project's existing dead-letter philosophy elsewhere in ingestion
    (gmail/sync.py's own parse-failure and fetch-failure handling): one
    corrupted, password-protected, or genuinely unsupported attachment is
    logged and skipped (returns ""), not allowed to lose the rest of the
    message's real body text or abort the sync. Filename extension is
    checked alongside mime_type since a real attachment's declared
    mimeType has been observed generic/wrong often enough elsewhere in
    this codebase (see message_parser.py's _strip_html real-testing
    notes) that trusting it alone would be fragile."""
    if len(data) > _MAX_ATTACHMENT_BYTES:
        return ""

    lower_name = filename.lower()
    try:
        if mime_type == _DOCX_MIME or lower_name.endswith(".docx"):
            return _extract_docx(data)
        if mime_type == _PDF_MIME or lower_name.endswith(".pdf"):
            return _extract_pdf(data)
        if mime_type in _IMAGE_MIMES or lower_name.endswith(_IMAGE_EXTENSIONS):
            return _extract_image(data)
    except Exception:  # noqa: BLE001 - deliberately broad, see docstring
        if logger is not None:
            logger.warning(
                f"attachment text extraction failed for {filename!r} ({mime_type})",
                extra={"operation": "gmail.attachment_extract", "status": "error", "duration_ms": 0},
            )
        return ""

    return ""


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(page for page in pages if page.strip())


def _extract_image(data: bytes) -> str:
    image = Image.open(io.BytesIO(data))
    return pytesseract.image_to_string(image).strip()
